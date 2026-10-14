"""STEP 1 - Loading. Only reads files; no cleaning here (see clean.py).

Run alone to test:   python -m src.load_data
"""
import re
import time

import pandas as pd
from openpyxl import load_workbook

from . import config as C


# ----------------------------------------------------------------------------
# File finding (works with "02 Field Production.xlsx" and "02_Field_Production.xlsx")
# ----------------------------------------------------------------------------
def find_file(prefix: str) -> "C.Path":
    """Find the Excel file starting with a 2-digit prefix, e.g. '02'.
    If duplicates like '..._1.xlsx' exist, the shortest name (the original) wins."""
    hits = [p for p in C.DATA_DIR.glob("*.xlsx") if re.match(rf"^{prefix}[ _]", p.name)]
    if not hits:
        raise FileNotFoundError(f"No file starting with '{prefix}' in {C.DATA_DIR.resolve()}. "
                               f"Set DWC_DATA_DIR or edit config.DATA_DIR.")
    return sorted(hits, key=lambda p: len(p.name))[0]


# ----------------------------------------------------------------------------
# 01 Introduction -> uncertainty parameters
# ----------------------------------------------------------------------------
def load_uncertainty() -> pd.DataFrame:
    """One row per case (index 1..100) with the 4 uncertainty parameters + split label."""
    wb = load_workbook(find_file("01"), read_only=True, data_only=True)
    rows = list(wb["Uncertainty Values"].iter_rows(values_only=True))
    wb.close()
    df = pd.DataFrame(rows[1:], columns=["case"] + C.PARAMS).dropna(subset=["case"])
    df["case"] = df["case"].astype(str).str.extract(r"(\d+)").astype(int)
    df = df.set_index("case").astype(float).sort_index()
    df["split"] = "train"
    df.loc[C.VAL_CASES, "split"] = "val"
    df.loc[C.TEST_CASES, "split"] = "test"
    return df


# ----------------------------------------------------------------------------
# 02 Field Production
# ----------------------------------------------------------------------------
def load_observed_table() -> pd.DataFrame:
    """The 'Field Production' sheet: observed quarterly rates incl. water injection."""
    wb = load_workbook(find_file("02"), read_only=True, data_only=True)
    rows = list(wb["Field Production"].iter_rows(values_only=True))
    wb.close()
    df = pd.DataFrame(rows[1:], columns=rows[0])
    df["Date"] = pd.to_datetime(df["Date"])
    return df.drop(columns=["Identifier"]).set_index("Date")


def _read_curve_sheet(wb, sheet: str) -> pd.DataFrame:
    """Sheet layout: row0 'Y1..', row1 'Date','Case_1'..'Case_100','RAMP_History', row2 units."""
    rows = list(wb[sheet].iter_rows(values_only=True))
    header = list(rows[1])
    df = pd.DataFrame(rows[3:], columns=header).rename(columns={"Date": "date"})
    df["date"] = pd.to_datetime(df["date"], errors="coerce")
    ren = {c: int(c.split("_")[1]) for c in df.columns if str(c).startswith("Case_")}
    ren["RAMP_History"] = "observed"
    df = df.rename(columns=ren)
    return df


def load_curve_sheets() -> dict:
    """Returns {'cum_gas': df, 'rate_gas': df, ...}; each df has 'date', 1..100, 'observed'."""
    wb = load_workbook(find_file("02"), read_only=True, data_only=True)
    out = {}
    for v in C.VARS:
        out[f"cum_{v}"] = _read_curve_sheet(wb, C.CUM_SHEETS[v])
        out[f"rate_{v}"] = _read_curve_sheet(wb, C.RATE_SHEETS[v])
    wb.close()
    return out


# ----------------------------------------------------------------------------
if __name__ == "__main__":
    t = time.time()
    U = load_uncertainty()
    print("Uncertainty table:", U.shape)
    print(U.groupby("split").size().to_dict())
    obs = load_observed_table()
    print("Observed table:", obs.shape, obs.index.min().date(), "->", obs.index.max().date())
    sheets = load_curve_sheets()
    for k, df in sheets.items():
        print(f"{k:9s} raw rows={len(df):6d}  rows with data={df[1].notna().sum():4d}  "
              f"observed pts={df['observed'].notna().sum()}")
    print(f"done in {time.time() - t:.1f}s")

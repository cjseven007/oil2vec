"""OPTIONAL - convert the huge 3-D grid workbooks (files 03-08, 250-330 MB each) to
compact per-case .npz files, once. Resumable: finished cases are skipped.

Reading one sheet takes ~10-15 s, so all 100 cases ~ 20-25 min (use --workers).
  python -m src.load_grids                    # all cases
  python -m src.load_grids --cases 1 2 3      # just a few (good for testing)
  python -m src.load_grids --workers 4

Each npz holds ACTIVE cells only (-999.25 sentinel rows removed), float32:
  data (n_active, 29), columns (names), full_shape_rows (rows incl. inactive)

NOTE (from the EDA): the dynamic columns (PRESSURE, SWAT, SOIL, ...) are identical in
every case = the INITIAL (1998) state, not the 2008 state. Only PORV/PORO/PERM*/TRAN*
change between cases.
"""
import argparse
import re
from concurrent.futures import ProcessPoolExecutor

import numpy as np
from openpyxl import load_workbook

from . import config as C
from . import load_data as L

SENTINEL = -999.25
GRID_PREFIXES = ["03", "04", "05", "06", "07", "08"]


def case_to_file_map() -> dict:
    """{case_number: (path, sheet_name)} by reading only the sheet names."""
    m = {}
    for pre in GRID_PREFIXES:
        try:
            path = L.find_file(pre)
        except FileNotFoundError:
            print(f"  (file {pre} not found - skipping)")
            continue
        wb = load_workbook(path, read_only=True)
        for sh in wb.sheetnames:
            m[int(re.search(r"\d+", sh).group())] = (path, sh)
        wb.close()
    return m


def convert_case(args):
    case, path, sheet = args
    out = C.CACHE_DIR / "grids" / f"case_{case:03d}.npz"
    if out.exists():
        return case, "skipped (exists)"
    wb = load_workbook(path, read_only=True, data_only=True)
    it = wb[sheet].iter_rows(values_only=True)
    header = [str(h).strip() for h in next(it)]
    arr = np.array(list(it), dtype=np.float64)
    wb.close()
    n_all = len(arr)
    active = arr[:, 7] != SENTINEL                     # PORO column == sentinel => inactive cell
    out.parent.mkdir(parents=True, exist_ok=True)
    np.savez_compressed(out, data=arr[active].astype(np.float32), columns=np.array(header),
                        full_shape_rows=n_all)
    return case, f"ok ({active.sum()} active of {n_all})"


def load_case_grid(case: int):
    """Use in your own code:  df = load_case_grid(5)  -> pandas DataFrame of active cells."""
    import pandas as pd
    z = np.load(C.CACHE_DIR / "grids" / f"case_{case:03d}.npz")
    return pd.DataFrame(z["data"], columns=list(z["columns"]))


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--cases", type=int, nargs="*", default=None)
    ap.add_argument("--workers", type=int, default=1)
    a = ap.parse_args()
    C.ensure_dirs()
    mp = case_to_file_map()
    cases = a.cases or sorted(mp)
    jobs = [(c, *mp[c]) for c in cases if c in mp]
    print(f"converting {len(jobs)} case(s) with {a.workers} worker(s)...")
    if a.workers > 1:
        with ProcessPoolExecutor(a.workers) as ex:
            for c, msg in ex.map(convert_case, jobs):
                print(f"case {c:3d}: {msg}")
    else:
        for j in jobs:
            c, msg = convert_case(j)
            print(f"case {c:3d}: {msg}")
